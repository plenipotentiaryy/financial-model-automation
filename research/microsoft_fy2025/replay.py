"""Reproduce a strict, report-only Microsoft FY2025 test from archived evidence.

This is a case-specific source adapter and independent audit, not a replacement
for the application's Excel calculation engine. It makes no network/AI calls.
"""

import argparse
import csv
import hashlib
import html
import json
import math
import re
import sys
from datetime import datetime, timezone
from decimal import Decimal, localcontext
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from finmodel.config import METADATA, NUMERIC_FIELDS, METRICS, METRIC_ORDER
from finmodel.outputs import ensure_available
from finmodel.schemas import Issue
from finmodel.service import build_company, configurations, load_dataset
from finmodel.template import set_value

CASE = Path(__file__).resolve().parent
URL = "https://www.microsoft.com/investor/reports/ar25/"
OUTPUT = ROOT / "outputs/microsoft_fy2025.xlsx"
D = Decimal


def archive_lines():
    lines, locations = {}, {}
    for path in sorted((CASE / "raw").glob("*.txt")):
        for physical_line, text in enumerate(path.read_text().splitlines(), 1):
            match = re.match(r"L(\d+): (.*)", text)
            if match:
                index, content = int(match[1]), match[2]
                if index in lines and lines[index] != content:
                    raise ValueError(f"Conflicting archived source line {index}")
                lines[index] = content
                locations[index] = {"file": str(path.relative_to(ROOT)), "line": physical_line, "web_extracted_line": index}
    return lines, locations


def number(text):
    text = text.replace("$", "").replace(",", "").strip()
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    return D(text)


def extract():
    lines, locators = archive_lines()
    rows = {year: {"company_id": "msft", "company_name": "Microsoft Corporation", "fiscal_year": year, "period_type": "FY", "scope": "consolidated", "currency": "USD", "monetary_unit": "millions", **{field: "" for field in NUMERIC_FIELDS}} for year in (2023, 2024, 2025)}
    ledger = {}

    def source(line, section):
        return {"url": URL, "section": section, "source_row": lines[line].split("|")[0].strip(), "archive_locator": locators[line]}

    def record(field, year, value, sources, transform="identity", note="", kind="reported"):
        rows[year][field] = str(value)
        ledger[(year, field)] = {"fiscal_year": year, "field": field, "input_value": str(value), "input_unit": "shares" if field == "shares_outstanding" else "USD millions", "upstream_status": kind, "transformation": transform, "sources": sources, "note": note}

    def table(line, years, expected_label):
        parts = [part.strip() for part in lines[line].split("|") if part.strip()]
        assert parts[0].startswith(expected_label), (line, parts[0])
        values = [number(value) for value in parts[1:]]
        assert len(values) == len(years), (line, parts)
        return dict(zip(years, values))

    years = (2025, 2024, 2023)
    statements = [("revenue", 814, "Total revenue"), ("cogs", 818, "Total cost of revenue"), ("gross_profit", 819, "Gross margin"), ("ebit", 823, "Operating income"), ("income_tax_expense", 826, "Provision for income taxes"), ("net_income", 827, "Net income")]
    for field, line, label in statements:
        for year, value in table(line, years, label).items():
            note = "Operating income is used as operating EBIT, not pretax income plus interest." if field == "ebit" else ("Microsoft's monetary Gross margin line is gross profit, not a percentage." if field == "gross_profit" else ("Total cost of revenue includes service/cloud costs; COGS proxy for this model." if field == "cogs" else ""))
            record(field, year, value, [source(line, "Income Statements")], note=note)
    for field, line, label in [("cash_flow_from_operations", 910, "Net cash from operations"), ("capex", 921, "Additions to property and equipment"), ("dividends_paid", 917, "Common stock cash dividends paid"), ("cash_and_equivalents", 931, "Cash and cash equivalents, end of period")]:
        for year, value in table(line, years, label).items():
            transform = "negate reported cash outflow" if field in {"capex", "dividends_paid"} else "identity"
            if field in {"capex", "dividends_paid"}:
                assert value < 0
                value = -value
            note = "Cash purchases of PP&E only; excludes noncash finance-lease additions, acquisitions and unpaid capex." if field == "capex" else ("Excludes short-term investments, even though Microsoft also reports combined liquidity." if field == "cash_and_equivalents" else "")
            record(field, year, value, [source(line, "Cash Flows Statements")], transform, note)
    for year, value in table(1132, years, "Interest expense").items():
        assert value < 0
        record("interest_expense", year, -value, [source(1132, "Note 3 — Other Income (Expense), Net")], "negate expense presentation", "Reported interest expense includes finance-lease interest and excludes capitalized interest; not cash interest paid.")
    for year, value in table(1734, years, "Balance, end of year").items():
        record("shares_outstanding", year, value * 1000000, [source(1734, "Note 15 — Shares Outstanding")], "millions of shares x 1,000,000", "FY-end shares, not weighted-average basic/diluted shares.")
    balances = [("total_assets", 864, "Total assets"), ("total_equity", 887, "Total stockholders"), ("current_assets", 857, "Total current assets"), ("current_liabilities", 874, "Total current liabilities"), ("inventory", 855, "Inventories"), ("accounts_receivable", 854, "Accounts receivable"), ("accounts_payable", 867, "Accounts payable")]
    for field, line, label in balances:
        for year, value in table(line, (2025, 2024), label).items():
            note = "Current net receivables only; long-term receivables and financing receivables are excluded." if field == "accounts_receivable" else ("Includes PP&E payable ($6.9bn FY2025/$4.3bn FY2024); DPO/CCC are coarse model proxies." if field == "accounts_payable" else "")
            record(field, year, value, [source(line, "Balance Sheets")], note=note)
    debt_lines = [(868, "Short-term debt"), (869, "Current portion of long-term debt"), (875, "Long-term debt")]
    debt_parts = [table(line, (2025, 2024), label) for line, label in debt_lines]
    for year in (2025, 2024):
        total = sum(parts[year] for parts in debt_parts)
        record("total_debt", year, total, [source(line, "Balance Sheets") for line, _ in debt_lines], "short-term debt + current long-term debt + noncurrent long-term debt", "Narrow borrowing-debt definition. Excludes finance and operating leases. Includes FY2024 commercial paper.", "derived_exact")
    assert "$40.4 billion" in lines[1527]
    record("market_value_debt", 2025, D("40400"), [source(1527, "Note 10 — Debt")], "40.4 billion USD x 1,000", "Reported estimate rounded to $0.1bn; long-term borrowings including current portion, with no commercial paper in FY2025; excludes leases.", "reported_rounded")

    gaps = {
        "ebitda": "No compatible reported EBITDA selected. Cash-flow D&A-and-other is not pure D&A. Notes 6/9 disclose rounded depreciation/amortization, so reconstruction would require an explicit definition and rounding policy; no estimate is inserted.",
        "monthly_recurring_revenue": "Company-wide fiscal-year-end MRR is not provided in the reviewed report. Cloud revenue, deferred revenue and remaining performance obligations are not MRR/ARR.",
        "market_cap": "No June 30 closing market capitalization in the selected financial statements. FY-end shares are available, but a verified matching-date share price is required.",
        "share_price": "No verified June 30 closing share price in the selected report. The total-return performance index is not a share price.",
        "pre_tax_cost_of_debt": "Current marginal market borrowing cost/YTM is not established. Interest expense / debt and lease discount rates are not substituted.",
        "risk_free_rate": "Requires a chosen currency-, date- and maturity-matched market assumption; not selected from the report.",
        "equity_beta": "Requires a documented estimate, benchmark and observation window; not selected from the report.",
        "equity_risk_premium": "Requires an explicit market-risk-premium assumption; not selected from the report.",
        "marginal_tax_rate": "Effective tax rate/provision is not a marginal tax-rate assumption for WACC.",
    }
    for year, row in rows.items():
        for field in NUMERIC_FIELDS:
            if (year, field) not in ledger:
                reason = gaps.get(field, "Not needed at this historical endpoint and not disclosed in the comparative balance sheets/notes selected for this case.")
                if field in gaps:
                    row[field] = "unavailable"
                rejected_sources = [source(896, "Cash Flows — mixed D&A line"), source(1386, "Note 6 — rounded depreciation"), source(1476, "Note 9 — rounded amortization")] if field == "ebitda" else ([source(1643, "Note 12 — unearned revenue, not ARR"), source(1654, "Note 12 — remaining obligations, not ARR")] if field == "monthly_recurring_revenue" else [])
                ledger[(year, field)] = {"fiscal_year": year, "field": field, "input_value": row[field], "input_unit": "", "upstream_status": "unavailable", "transformation": None, "sources": rejected_sources, "note": reason}

    # Independent source reconciliations (USD millions), before any model build.
    checks = []
    for year in years:
        row = rows[year]
        assert D(row["revenue"]) - D(row["cogs"]) == D(row["gross_profit"])
        opex = sum(table(line, years, label)[year] for line, label in [(820, "Research and development"), (821, "Sales and marketing"), (822, "General and administrative")])
        assert D(row["gross_profit"]) - opex == D(row["ebit"])
        assert table(825, years, "Income before income taxes")[year] - D(row["income_tax_expense"]) == D(row["net_income"])
        checks.append({"year": year, "income_statement_reconciliations": "PASS"})
    for year in (2024, 2025):
        row = rows[year]
        assert table(881, (2025, 2024), "Total liabilities")[year] + D(row["total_equity"]) == D(row["total_assets"])
        assert table(1522, (2025, 2024), "Total debt")[year] + debt_parts[0][year] == D(row["total_debt"])
        assert table(851, (2025, 2024), "Cash and cash equivalents")[year] == D(row["cash_and_equivalents"])
        checks.append({"year": year, "balance_sheet_and_debt_reconciliations": "PASS"})
    return [rows[year] for year in sorted(rows)], list(ledger.values()), gaps, checks


def expected_metrics(rows):
    """Independent case audit, with Decimal; never used to populate the workbook."""
    by_year = {row["fiscal_year"]: row for row in rows}
    c, p, first = by_year[2025], by_year[2024], by_year[2023]
    val = lambda row, field: D(row[field])
    avg = lambda field: (val(c, field) + val(p, field)) / 2
    revenue, ebit, income = val(c, "revenue"), val(c, "ebit"), val(c, "net_income")
    fcf = val(c, "cash_flow_from_operations") - val(c, "capex")
    dso = avg("accounts_receivable") / revenue * 365
    dio = avg("inventory") / val(c, "cogs") * 365
    dpo = avg("accounts_payable") / val(c, "cogs") * 365
    return {
        "revenue_growth_yoy": revenue / val(p, "revenue") - 1,
        "net_debt": (val(c, "total_debt") - val(c, "cash_and_equivalents")) * 1000000,
        "gross_margin": val(c, "gross_profit") / revenue,
        "ebit_margin": ebit / revenue, "net_profit_margin": income / revenue,
        "free_cash_flow": fcf * 1000000, "fcf_margin": fcf / revenue,
        "current_ratio": val(c, "current_assets") / val(c, "current_liabilities"),
        "quick_ratio": (val(c, "current_assets") - val(c, "inventory")) / val(c, "current_liabilities"),
        "interest_coverage": ebit / val(c, "interest_expense"),
        "roa": income / avg("total_assets"), "roe": income / avg("total_equity"),
        "revenue_cagr": (revenue / val(first, "revenue")).sqrt() - 1,
        "dso": dso, "dio": dio, "dpo": dpo, "cash_conversion_cycle": dso + dio - dpo,
    }


def notes_for(metric_id):
    notes = {
        "net_debt": "Borrowing debt less cash equivalents only; excludes $64.323bn short-term investments and $46.172bn finance-lease liabilities.",
        "ebit_margin": "Operating income is the selected operating-EBIT definition.",
        "interest_coverage": "Operating income / reported interest expense; excludes capitalized interest from the denominator.",
        "free_cash_flow": "CFO less cash PP&E additions. Noncash finance-lease additions and acquisitions are excluded.",
        "fcf_margin": "Uses cash-PP&E FCF, not lease-adjusted FCFF/FCFE.",
        "revenue_growth_yoy": "Reported consolidated growth; acquisition effects are included, not organic growth.",
        "revenue_cagr": "FY2023–FY2025, two elapsed years; acquisition effects included.",
        "dso": "Proxy: average current net receivables / total revenue x 365; credit sales unavailable.",
        "dio": "Proxy: average inventory / total cost of revenue x 365; service-heavy denominator.",
        "dpo": "Coarse proxy: average total AP / cost of revenue x 365. AP includes capex; credit purchases unavailable.",
        "cash_conversion_cycle": "Arithmetic CCC under model proxies. Capex AP and service revenue limit economic comparability; negative does not prove superior cash management.",
        "roa": "Two-point average total assets at June 30, 2024 and 2025.",
        "roe": "Two-point average stockholders' equity at June 30, 2024 and 2025.",
    }
    return notes.get(metric_id, "Computed from compatible reported financial-statement inputs under the documented model convention.")


def case_template(ledger, target):
    wb = load_workbook(ROOT / "samples/template.xlsx")
    ws = wb.create_sheet("Filing_Sources")
    ws.append(["Microsoft Annual Report 2025 — upstream source ledger"])
    ws.append([URL])
    headers = ["FY", "Field", "Input value", "Unit", "Upstream status", "Transformation", "Source section / row", "Saved evidence locator", "Notes"]
    ws.append(headers)
    for entry in ledger:
        loc = "; ".join(f"{s['archive_locator']['file']}:{s['archive_locator']['line']} (web L{s['archive_locator']['web_extracted_line']})" for s in entry["sources"])
        source = "; ".join(f"{s['section']} / {s['source_row']}" for s in entry["sources"])
        values = [entry["fiscal_year"], entry["field"], entry["input_value"], entry["input_unit"], entry["upstream_status"], entry["transformation"], source, loc, entry["note"]]
        for col, value in enumerate(values, 1):
            set_value(ws.cell(ws.max_row + 1 if col == 1 else ws.max_row, col), value)
    for col, width in zip("ABCDEFGHI", [10, 29, 21, 18, 23, 40, 65, 80, 95]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:I{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for row in ws:
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row[0].row].height = 78 if row[0].row > 3 else 24
    for cell in ws[3]:
        cell.fill = PatternFill("solid", fgColor="253E5E")
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
    ws["A1"].alignment = Alignment(vertical="center", wrap_text=False)
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="253E5E")
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=False)
    ws["A2"].hyperlink = URL
    ws["A2"].font = Font(name="Arial", size=10, color="1C5BBF", underline="single")
    report = wb["Report"]
    report["A44"] = "Real-data case: Microsoft FY2025, June 30 year end. Strict report-only inputs; see Filing_Sources."
    report["A45"] = "EBITDA is unresolved. Missing market/recurring-revenue assumptions are not estimated."
    report["A46"] = "Debt/cash and working-capital proxies have material scope limitations; see case report."
    for row in range(7, 35):
        report.row_dimensions[row].height = 60  # Fit source-specific caveats without truncation.
    wb.save(target)


def main(overwrite=False):
    input_path, ledger_path, template = CASE / "financials.csv", CASE / "source_ledger.json", CASE / "microsoft_template.xlsx"
    ensure_available([input_path, ledger_path, template, OUTPUT, OUTPUT.with_name(OUTPUT.stem + ".reports")], overwrite)
    rows, ledger, gaps, source_checks = extract()
    with input_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=METADATA + NUMERIC_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    ledger_path.write_text(json.dumps(ledger, indent=2, ensure_ascii=False) + "\n")
    case_template(ledger, template)
    config = configurations(ROOT / "config/aliases.yaml", ROOT / "config/template_map.yaml", ROOT / "config/metric_definitions.yaml")
    company = load_dataset([input_path], config).companies["msft"]
    for metric_id, entry in company.metric_statuses.items():
        if entry["status"] == "OK":
            entry["message"] = notes_for(metric_id)
    company.issues.append(Issue("SOURCE_SCOPE_LIMITATIONS", "Report-only test. EBITDA and market assumptions are unresolved. Net debt excludes investments and leases; CCC uses coarse proxies. See Filing_Sources and the case audit.", company_id="msft", source=URL))
    manifest = build_company(company, template, OUTPUT, config, overwrite=overwrite, calculate=True)
    assert manifest["outcome"] != "FAILED", manifest["issues"]
    actual = load_workbook(OUTPUT, data_only=True)
    formula_wb = load_workbook(OUTPUT, data_only=False)
    baseline = load_workbook(template, data_only=False)
    for ws in baseline:
        for row in ws:
            for cell in row:
                if cell.data_type == "f":
                    assert formula_wb[ws.title][cell.coordinate].value == cell.value
    with localcontext() as context:
        context.prec = 40
        expected = expected_metrics(rows)
    audit = []
    for metric_id in METRIC_ORDER:
        definition = METRICS[metric_id]
        sheet, address = definition["cell"].split("!")
        r_sheet, r_address = definition["report_cell"].split("!")
        value, report_value = actual[sheet][address].value, actual[r_sheet][r_address].value
        status = company.metric_statuses[metric_id]["status"]
        if metric_id in expected:
            assert status == "OK"
            assert math.isclose(value, float(expected[metric_id]), rel_tol=1e-12, abs_tol=1e-8), (metric_id, value, expected[metric_id])
            assert value == report_value
            note = notes_for(metric_id)
        else:
            assert status == "MISSING_INPUTS", (metric_id, status)
            assert value in {None, ""} and report_value in {None, ""}
            missing = [req.split(":")[0] for req in definition["requires"] if company.value(req.split(":")[0], company.period_year(req.split(":")[1])) is None]
            note = " ".join(gaps.get(field, "Unavailable input: " + field) for field in dict.fromkeys(missing))
        audit.append({"metric_id": metric_id, "label": definition["label"], "status": status, "value": value if value != "" else None, "unit": definition["unit"], "expected_independent": str(expected[metric_id]) if metric_id in expected else None, "verification": "PASS" if metric_id in expected else "PASS_UNAVAILABLE", "interpretation": note})
    assert len(expected) == 17 and len(audit) == 28
    source_files = [{"path": str(path.relative_to(ROOT)), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for path in sorted((CASE / "raw").glob("*.txt"))]
    evidence = {"company": "Microsoft Corporation", "fiscal_year_end": "2025-06-30", "historical_years": [2023, 2024, 2025], "basis": "Historical FY2025 report-only test, not a current valuation or latest-report claim", "source_url": URL, "retrieved_on": "2026-10-15", "source_acquisition": "Official report excerpts retrieved through web tool. Direct urllib HTML download returned HTTP 403; original HTML was not archived.", "source_files": source_files, "metrics_computed": 17, "metrics_unavailable": 11, "all_28_verified": True, "source_reconciliations": source_checks, "calculation": manifest["calculation"], "source_gaps": gaps, "audit": audit}
    (CASE / "verification.json").write_text(json.dumps(evidence, indent=2, ensure_ascii=False) + "\n")
    (CASE / "raw/source_manifest.json").write_text(json.dumps({key: evidence[key] for key in ("source_url", "retrieved_on", "source_acquisition", "source_files")}, indent=2) + "\n")
    table = []
    for item in audit:
        value = item["value"]
        display = "—" if value is None else (f"{value * 100:.4f}%" if item["unit"] == "%" else (f"{value / 1000000:,.3f} млн USD" if item["unit"] == "currency" else f"{value:.4f} {item['unit']}"))
        table.append("<tr>" + "".join(f"<td>{html.escape(str(x))}</td>" for x in [item["label"], display, item["verification"], item["interpretation"]]) + "</tr>")
    report_html = '''<!doctype html><meta charset="utf-8"><title>Microsoft FY2025 — проверка 28 метрик</title><style>body{font:16px/1.5 system-ui;max-width:1500px;margin:40px auto;padding:0 24px;color:#203047}h1{font-size:28px}table{border-collapse:collapse;width:100%}th,td{text-align:left;padding:12px;border-bottom:1px solid #dde3ea;vertical-align:top}th{background:#253e5e;color:white}td:nth-child(2){white-space:nowrap}a{color:#1c5bbf}.note{padding:18px;background:#fff4d9}</style>
<h1>Microsoft FY2025: 17 рассчитано, 11 недоступно</h1>
<p>Год закончился 30 июня 2025. История FY2023–FY2025. Все 28 позиций проверены: рассчитанные результаты сверены независимой арифметикой, недоступные — с пустыми Excel-результатами. Формулы пересчитаны LibreOffice.</p>
<p><a href="https://www.microsoft.com/investor/reports/ar25/">Официальный годовой отчёт Microsoft</a> · <a href="../../outputs/microsoft_fy2025.xlsx">Заполненная модель</a> · <a href="source_ledger.json">Источники каждого входа</a> · <a href="verification.json">Машинная проверка всех 28</a></p>
<div class="note"><b>Границы проверки.</b> Это исторический тест по выбранной отчётности, а не текущая оценка компании. EBITDA не подменена смешанной строкой D&amp;A and other. Рыночные цены, MRR и предпосылки WACC не придуманы. Net debt исключает short-term investments и аренду. DPO/CCC — приближённые показатели: кредиторская задолженность Microsoft включает capex, а знаменатель — cost of revenue. Расчёт по формуле не доказывает экономическую сопоставимость.</div>
<table><thead><tr><th>Метрика</th><th>Результат</th><th>Проверка</th><th>Определение / чего не хватает</th></tr></thead><tbody>''' + "\n".join(table) + "</tbody></table>"
    (CASE / "report.html").write_text(report_html, encoding="utf-8")
    print(json.dumps({"computed": 17, "unavailable": 11, "all_28_verified": True, "workbook": str(OUTPUT), "audit": str(CASE / "report.html")}, indent=2))
    for item in audit:
        print(item["metric_id"], item["value"], item["status"])


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--overwrite", action="store_true")
    main(parser.parse_args().overwrite)
