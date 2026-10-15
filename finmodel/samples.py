"""Self-contained fixtures; generated financial data are explicitly fictional."""

import csv
import json
from pathlib import Path

import yaml
from openpyxl import Workbook

from .config import ALIASES, METADATA, METRICS, TEMPLATE_MAP, NUMERIC_FIELDS
from .outputs import ensure_available
from .template import create_template, set_value

HEADERS = METADATA + ["revenue", "ebitda", "total_debt", "cash_and_equivalents"]
ACME = [
    ["acme", "Acme Example", 2024, "FY", "consolidated", "PLN", "millions", 80, None, None, None],
    ["acme", "Acme Example", 2025, "FY", "consolidated", "PLN", "millions", 100, 20, 35, 5],
]

# Independently specified fictional fixtures, not a production calculation engine.
FULL_DATA = []
for year, values in [
    (2023, {"revenue": 64, "ebitda": 12.8, "total_debt": 25, "cash_and_equivalents": 3, "cogs": 42, "gross_profit": 22, "ebit": 10, "net_income": 6, "interest_expense": 2, "income_tax_expense": 2, "cash_flow_from_operations": 11, "capex": 4, "dividends_paid": 1, "total_assets": 90, "total_equity": 35, "current_assets": 40, "current_liabilities": 20, "inventory": 8, "accounts_receivable": 6, "accounts_payable": 6, "market_cap": 100, "share_price": 10, "shares_outstanding": 10000000, "monthly_recurring_revenue": 1.2, "market_value_debt": 35}),
    (2024, {"revenue": 80, "ebitda": 16, "total_debt": 30, "cash_and_equivalents": 4, "cogs": 50, "gross_profit": 30, "ebit": 12, "net_income": 8, "interest_expense": 2, "income_tax_expense": 2, "cash_flow_from_operations": 14, "capex": 5, "dividends_paid": 1.5, "total_assets": 100, "total_equity": 40, "current_assets": 50, "current_liabilities": 25, "inventory": 10, "accounts_receivable": 8, "accounts_payable": 8, "market_cap": 120, "share_price": 12, "shares_outstanding": 10000000, "monthly_recurring_revenue": 1.6, "market_value_debt": 40}),
    (2025, {"revenue": 100, "ebitda": 20, "total_debt": 35, "cash_and_equivalents": 5, "cogs": 60, "gross_profit": 40, "ebit": 15, "net_income": 10, "interest_expense": 3, "income_tax_expense": 2, "cash_flow_from_operations": 18, "capex": 6, "dividends_paid": 2, "total_assets": 120, "total_equity": 50, "current_assets": 60, "current_liabilities": 30, "inventory": 15, "accounts_receivable": 12, "accounts_payable": 12, "market_cap": 150, "share_price": 15, "shares_outstanding": 10000000, "monthly_recurring_revenue": 2, "market_value_debt": 50}),
]:
    FULL_DATA.append({"company_id": "complete", "company_name": "Complete Example (fictional)", "fiscal_year": year, "period_type": "FY", "scope": "consolidated", "currency": "PLN", "monetary_unit": "millions", **values, "pre_tax_cost_of_debt": 0.06, "risk_free_rate": 0.04, "equity_beta": 1.2, "equity_risk_premium": 0.05, "marginal_tax_rate": 0.2})

FULL_EXPECTED = {"revenue_growth_yoy": 0.25, "ebitda_margin": 0.2, "net_debt": 30000000, "net_debt_to_ebitda": 1.5, "arr": 24000000, "wacc": 0.087, "gross_margin": 0.4, "ebit_margin": 0.15, "net_profit_margin": 0.1, "free_cash_flow": 12000000, "fcf_margin": 0.12, "current_ratio": 2, "quick_ratio": 1.5, "interest_coverage": 5, "roa": 0.09090909090909091, "roe": 0.2222222222222222, "enterprise_value": 180000000, "ev_to_revenue": 1.8, "ev_to_ebitda": 9, "price_to_earnings": 15, "revenue_cagr": 0.25, "ebitda_cagr": 0.25, "net_leverage_trend": -0.125, "cash_conversion": 0.9, "dso": 36.5, "dio": 76.04166666666667, "dpo": 60.833333333333336, "cash_conversion_cycle": 51.708333333333336}


def write_full_samples(directory):
    headers = METADATA + NUMERIC_FIELDS
    with (directory / "all_metrics.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(FULL_DATA)
    (directory / "all_metrics.json").write_text(json.dumps(FULL_DATA, indent=2) + "\n", encoding="utf-8")
    (directory / "all_metrics_expected.json").write_text(json.dumps(FULL_EXPECTED, indent=2) + "\n", encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "Financials"
    ws.append(headers)
    for entry in FULL_DATA:
        ws.append([entry.get(field) for field in headers])
    wb.save(directory / "all_metrics.xlsx")


def write_csv(path, rows):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(HEADERS)
        writer.writerows(rows)


def initialize(root: Path, overwrite=False):
    names = ["config/aliases.yaml", "config/template_map.yaml", "config/metric_definitions.yaml", "samples/template.xlsx", "samples/acme.csv", "samples/acme.json", "samples/acme.xlsx", "samples/missing_inputs.csv", "samples/invalid_periods.csv", "samples/companies/acme.csv", "samples/companies/beta.csv", "samples/all_metrics.csv", "samples/all_metrics.json", "samples/all_metrics.xlsx", "samples/all_metrics_expected.json"]
    ensure_available([root / name for name in names], overwrite)
    for name in names:
        (root / name).parent.mkdir(parents=True, exist_ok=True)
    for name, value in [("aliases.yaml", ALIASES), ("template_map.yaml", TEMPLATE_MAP), ("metric_definitions.yaml", METRICS)]:
        (root / "config" / name).write_text(yaml.safe_dump(value, sort_keys=False, allow_unicode=True), encoding="utf-8")
    create_template(root / "samples/template.xlsx")
    write_csv(root / "samples/acme.csv", ACME)
    (root / "samples/acme.json").write_text(json.dumps([dict(zip(HEADERS, row)) for row in ACME], indent=2) + "\n", encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "Financials"
    for number, row in enumerate([HEADERS, *ACME], 1):
        for column, value in enumerate(row, 1):
            set_value(ws.cell(number, column), value)
    wb.save(root / "samples/acme.xlsx")
    missing = [list(row) for row in ACME]
    missing[1][8] = None
    write_csv(root / "samples/missing_inputs.csv", missing)
    invalid = [list(row) for row in ACME]
    invalid[1][3] = "TTM"
    write_csv(root / "samples/invalid_periods.csv", invalid)
    write_csv(root / "samples/companies/acme.csv", ACME)
    beta = [["beta", "Beta Example", 2024, "FY", "consolidated", "EUR", "thousands", 200, None, None, None], ["beta", "Beta Example", 2025, "FY", "consolidated", "EUR", "thousands", 300, 60, 40, 10]]
    write_csv(root / "samples/companies/beta.csv", beta)
    write_full_samples(root / "samples")
    return names
