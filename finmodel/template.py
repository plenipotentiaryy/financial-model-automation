"""Example template creation and strict write-allowlist validation."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.properties import CalcProperties

from .config import METRICS, METRIC_ORDER, NUMERIC_FIELDS, TEMPLATE_MAP, RATES, REPORT_END_ROW
from .schemas import ModelError

SHEETS = ["Report", "Assumptions", "Calculations", "Inputs", "Checks", "Source_Map", "Data_Dictionary"]
MONEY_FORMAT = '#,##0.0,,;(#,##0.0,,);"—"'
NUMBER_FORMAT = '#,##0.0;(#,##0.0);"—"'
PERCENT_FORMAT = '0.0%;(0.0%);"—"'
LEVERAGE_FORMAT = '0.0"x";(0.0"x");"—"'


def set_value(cell, value):
    """Untrusted strings are literal inline strings, never spreadsheet formulas."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def request_recalculation(wb):
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)


def create_template(path: Path, metrics=None):
    metrics = metrics or METRICS
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEETS:
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        for col in "ABCDEFG":
            ws.column_dimensions[col].width = 24
        ws.column_dimensions["A"].width = 38
        ws.print_options.horizontalCentered = False
    titles = {"Report": "Financial KPI report — 28 metrics", "Assumptions": "Explicit cost-of-capital assumptions", "Checks": "Validation issues", "Source_Map": "Source-to-cell lineage", "Data_Dictionary": "Metric definitions"}
    for name, title in titles.items():
        wb[name]["A2"] = title
    wb["Data_Dictionary"]["A1"] = "Template version"
    wb["Data_Dictionary"]["B1"] = 2
    inputs = wb["Inputs"]
    for address, value in {"A2": "Company", "A3": "Fiscal year", "A4": "Currency", "A5": "Company ID / scope", "A6": "Amounts shown in millions", "B6": "Previous FY", "C6": "Latest FY", "D6": "First FY (CAGR only)"}.items():
        inputs[address] = value
    inputs["B4"] = '=IF(C4<>"",C4,"")'
    labels = {"ebitda": "EBITDA", "ebit": "EBIT", "cogs": "Cost of goods sold", "cash_flow_from_operations": "Operating cash flow", "capex": "Capex (positive spending)", "monthly_recurring_revenue": "Monthly recurring revenue (MRR)", "market_value_debt": "Market value of debt (millions)"}
    for group in ("financials", "assumptions"):
        for field, periods in TEMPLATE_MAP[group].items():
            for period, address in periods.items():
                name, cell, bounds = split_address(address)
                ws = wb[name]
                row = bounds[1]
                ws.cell(row, 1, labels.get(field, field.replace("_", " ").capitalize()))
                target = ws[cell]
                target.font = Font(name="Arial", size=10, color="0000FF")
                target.fill = PatternFill("solid", fgColor="EFF4FC")
                target.number_format = PERCENT_FORMAT if field in RATES else ('0.00' if field in {"equity_beta", "share_price"} else ('#,##0' if field == "shares_outstanding" else MONEY_FORMAT))
    input_end = inputs.max_row
    inputs.cell(input_end + 2, 1, "Flows: full FY. Balances and market values: same FY end.")
    inputs.cell(input_end + 3, 1, "Fiscal year labels do not imply 31 December.")
    inputs.cell(input_end + 4, 1, "Share price is per share; shares are counts. Other money is shown in millions.")
    inputs.freeze_panes = "B7"
    assumptions = wb["Assumptions"]
    assumptions["A4"], assumptions["B4"] = "Assumption", "Explicit FY-end value"
    assumptions["A12"] = "Rates are fractions (0.04 = 4%). No assumptions are estimated."
    assumptions["A13"] = "Market debt is explicit; reported/book debt is not substituted."
    assumptions["A14"] = "WACC: two capital components, CAPM cost of equity."
    calculations = wb["Calculations"]
    calculations["A1"] = "Template-owned formulas. Money displayed in millions; days use a 365-day convention."
    formats = {"%": PERCENT_FORMAT, "currency": MONEY_FORMAT, "x": LEVERAGE_FORMAT, "days": NUMBER_FORMAT}
    report = wb["Report"]
    report["A3"], report["A4"] = "Company", "Fiscal year / currency"
    report["B3"] = '=IF(Inputs!B2<>"",Inputs!B2,"")'
    report["B4"] = '=IF(ISNUMBER(Inputs!C3),Inputs!C3,"")'
    report["C4"] = '=IF(Inputs!C4<>"",Inputs!C4,"")'
    for col, label in enumerate(["Metric", "Result", "Basis / units", "Input status at build", "Details at build"], 1):
        report.cell(6, col, label)
    for index, metric_id in enumerate(METRIC_ORDER):
        definition = metrics[metric_id]
        row = index + 2
        calculations.cell(row, 1, definition["label"])
        calculations.cell(row, 2, definition["excel_formula"])
        calculations.cell(row, 2).number_format = formats[definition["unit"]]
        report_row = index + 7
        report.cell(report_row, 1, definition["label"])
        report.cell(report_row, 2, f'=IF(ISNUMBER(Calculations!B{row}),Calculations!B{row},"")')
        report.cell(report_row, 2).number_format = formats[definition["unit"]]
        if definition["unit"] == "currency":
            report.cell(report_row, 3, '=IF(Inputs!C4<>"",Inputs!C4&", millions","")')
        else:
            report.cell(report_row, 3, {"%": "Annual / FY-end basis", "x": "Multiple", "days": "365-day convention"}[definition["unit"]])
    report["C12"] = "Market weights / CAPM"
    report["C29"] = "Change in turns"
    calculations.cell(len(metrics) + 4, 1, "Unavailable inputs or invalid denominators leave results blank; see Checks.")
    report["A37"] = "Calculation state"
    report["A40"] = "Statuses describe imported data at build time. Rerun after editing inputs."
    report["A41"] = "CAGR: first to latest FY. ROA/ROE: average balances. Cash conversion: CFO / EBITDA."
    report["A42"] = "CCC: DSO + DIO - DPO; revenue and COGS are documented credit-sales/purchases proxies."
    report.freeze_panes = "B7"
    headers = {
        "Checks": ["Severity", "Code", "Metric", "Company ID", "Message", "Source"],
        "Source_Map": ["Source file", "Source sheet", "Source row", "Source column", "Locator", "Field", "Fiscal year", "Period", "Scope", "Currency", "Original value", "Original unit", "Normalized value", "Canonical unit", "Value status", "Destination", "Dependent metrics"],
        "Data_Dictionary": ["Metric ID", "Label", "Formula / definition", "Required fields", "Unit", "Implementation", "Limitations"],
    }
    for name, values in headers.items():
        ws = wb[name]
        for col, label in enumerate(values, 1):
            ws.cell(4, col, label)
            if name in {"Checks", "Source_Map"}:
                ws.cell(5, col).alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "C5" if name == "Source_Map" else "A5"
        end_row = 10004 if name == "Source_Map" else (1004 if name == "Checks" else 4 + len(metrics))
        ws.auto_filter.ref = f"A4:{ws.cell(4, len(values)).column_letter}{end_row}"
    for row, metric_id in enumerate(METRIC_ORDER, 5):
        definition = metrics[metric_id]
        values = [metric_id, definition["label"], definition["formula"], ", ".join(definition["requires"]), definition["unit"], definition["implementation"], definition["limitations"]]
        for col, value in enumerate(values, 1):
            set_value(wb["Data_Dictionary"].cell(row, col), value)
    for name in SHEETS:
        ws = wb[name]
        for row in ws:
            for cell in row:
                if cell.font.name != "Arial":
                    cell.font = Font(name="Arial", size=10, color="202B3C")
                cell.alignment = Alignment(vertical="center", wrap_text=name in headers)
                if cell.data_type == "f" and name != "Report":
                    cell.font = Font(name="Arial", size=10, color="008000" if "!" in cell.value else "000000")
        if name in titles:
            ws["A2"].font = Font(name="Arial", size=14, bold=True, color="1E3554")
            ws["A2"].alignment = Alignment(vertical="center", wrap_text=False)
        header_row = 6 if name in {"Inputs", "Report"} else (4 if name in headers or name == "Assumptions" else None)
        if header_row:
            end = 5 if name == "Report" else (4 if name == "Inputs" else (2 if name == "Assumptions" else len(headers[name])))
            for cell in ws[header_row][:end]:
                cell.fill = PatternFill("solid", fgColor="253E5E")
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[header_row].height = 30
        if name in headers:
            for row in range(5, ws.max_row + 1):
                ws.row_dimensions[row].height = 72 if name == "Data_Dictionary" else 66
        if name == "Report":
            for row in range(7, REPORT_END_ROW + 1):
                ws.row_dimensions[row].height = 48
                for col in (1, 3, 4, 5):
                    ws.cell(row, col).alignment = Alignment(vertical="center", wrap_text=True)
                ws.cell(row, 2).font = Font(name="Arial", size=10, color="202B3C")
                ws.cell(row, 2).border = Border(bottom=Side(style="hair", color="DDE4ED"))
    report.column_dimensions["D"].width = 28
    report.column_dimensions["E"].width = 82
    wb["Report"].sheet_properties.tabColor = "253E5E"
    wb["Assumptions"].sheet_properties.tabColor = "7391B2"
    for name, widths in {"Checks": {"A": 12, "B": 30, "C": 30, "D": 20, "E": 90, "F": 55}, "Source_Map": {"A": 58, "B": 20, "C": 14, "D": 25, "E": 35, "F": 28, "G": 14, "H": 12, "I": 20, "J": 14, "K": 22, "L": 18, "M": 24, "N": 20, "O": 20, "P": 36, "Q": 64}, "Data_Dictionary": {"A": 28, "B": 36, "C": 66, "D": 80, "E": 14, "F": 20, "G": 90}}.items():
        for col, width in widths.items():
            wb[name].column_dimensions[col].width = width
    wb["Report"].conditional_formatting.add(f"D7:D{REPORT_END_ROW}", FormulaRule(formula=['D7<>"OK"'], fill=PatternFill("solid", fgColor="FFF0D2")))
    wb["Checks"].conditional_formatting.add("A5:F1004", FormulaRule(formula=['$A5="error"'], fill=PatternFill("solid", fgColor="FFE0E0")))
    request_recalculation(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def split_address(address: str):
    try:
        sheet, target = address.rsplit("!", 1)
        sheet = sheet.strip("'")
        bounds = range_boundaries(target)
        if not all(isinstance(v, int) for v in bounds):
            raise ValueError("Unbounded range")
        if not (1 <= bounds[0] <= bounds[2] <= 16384 and 1 <= bounds[1] <= bounds[3] <= 1048576):
            raise ValueError("Range exceeds worksheet")
        return sheet, target, bounds
    except (ValueError, AttributeError, TypeError) as exc:
        raise ModelError(f"Invalid cell/range address: {address!r}") from exc


def validate_map(wb, mapping):
    if not isinstance(mapping, dict) or any(not isinstance(mapping.get(key), dict) for key in ("metadata", "financials", "assumptions", "regions")):
        raise ModelError("Template map sections must be mappings")
    if mapping.get("version") != 2:
        raise ModelError("Template map v2 is required for the 28-metric template")
    if set(mapping) != {"version", "metadata", "financials", "assumptions", "regions"}:
        raise ModelError("Template map must define version, metadata, financials, assumptions and regions")
    if set(mapping["metadata"]) != set(TEMPLATE_MAP["metadata"]):
        raise ModelError("Template map metadata keys must match the documented contract")
    if set(mapping["regions"]) != set(TEMPLATE_MAP["regions"]):
        raise ModelError("Missing/unknown diagnostic region")
    addresses = []
    mapped_observations = set()
    for name, address in mapping["metadata"].items():
        addresses.append((name, address, True))
    for kind in ("financials", "assumptions"):
        for field, periods in mapping[kind].items():
            if field not in NUMERIC_FIELDS or not isinstance(periods, dict) or not periods or set(periods) - {"previous", "latest", "start"}:
                raise ModelError(f"Invalid {kind} mapping: {field}")
            for period, address in periods.items():
                if (field, period) in mapped_observations:
                    raise ModelError(f"Field/period mapped more than once: {field}:{period}")
                mapped_observations.add((field, period))
                addresses.append((f"{field}:{period}", address, True))
    for name, address in mapping["regions"].items():
        addresses.append((name, address, False))
    rectangles = []
    for name, address, scalar in addresses:
        sheet, _, bounds = split_address(address)
        if sheet not in wb.sheetnames:
            raise ModelError(f"Map references absent sheet: {sheet}")
        x1, y1, x2, y2 = bounds
        if scalar and (x1 != x2 or y1 != y2):
            raise ModelError(f"{name} must map to one cell")
        if (x2 - x1 + 1) * (y2 - y1 + 1) > 250000:
            raise ModelError(f"Write region too large: {name}")
        for other_sheet, (a1, b1, a2, b2) in rectangles:
            if sheet == other_sheet and x1 <= a2 and a1 <= x2 and y1 <= b2 and b1 <= y2:
                raise ModelError(f"Overlapping write destinations: {name}")
        rectangles.append((sheet, bounds))
        ws = wb[sheet]
        for merged in ws.merged_cells.ranges:
            a1, b1, a2, b2 = merged.bounds
            if x1 <= a2 and a1 <= x2 and y1 <= b2 and b1 <= y2:
                raise ModelError(f"Write region intersects merged cells: {name}")
        # Inspect instantiated cells only; do not materialize huge empty regions.
        for row in ws.iter_rows(min_row=y1, max_row=min(y2, ws.max_row), min_col=x1, max_col=min(x2, ws.max_column)):
            if any(cell.data_type == "f" for cell in row):
                raise ModelError(f"Write destination intersects a formula: {name}")
    for group in ("financials", "assumptions"):
        for metric, periods in TEMPLATE_MAP[group].items():
            if not set(periods) <= set(mapping[group].get(metric, {})):
                raise ModelError(f"Required {group} mapping absent: {metric}")
    return rectangles


def load_template(path: Path, mapping):
    if path.suffix.lower() != ".xlsx":
        raise ModelError("Only .xlsx templates are supported")
    wb = load_workbook(path, data_only=False, keep_links=True)
    if set(SHEETS) - set(wb.sheetnames):
        raise ModelError("Template is missing required worksheets")
    if wb.vba_archive or wb._external_links:
        raise ModelError("Macros and external workbook links are not supported")
    if wb["Data_Dictionary"]["B1"].value != 2:
        raise ModelError("Template/map version mismatch: a v2 workbook is required for all 28 metrics")
    for metric_id, definition in METRICS.items():
        for key in ("cell", "report_cell"):
            sheet, address, _ = split_address(definition[key])
            if wb[sheet][address].data_type != "f":
                raise ModelError(f"Template is missing the {key} formula for {metric_id}")
    validate_map(wb, mapping)
    return wb
