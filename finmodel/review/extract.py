"""Conservative extraction: preserve labels, positions and unknowns, never fuzzy-map."""
import csv
import hashlib
import json
import re
from pathlib import Path

from ..config import ALIASES, METADATA
from ..mapping import alias_lookup, header_key
from ..schemas import ModelError
from .contracts import Candidate
from .store import Store

# Context-sensitive financial aliases. Gross margin is deliberately NOT an alias:
# it may mean a currency subtotal or a percentage.
LABELS = {'operating income': 'ebit', 'net income': 'net_income', 'total revenue': 'revenue',
          'total revenues': 'revenue', 'cost of revenue': 'cogs', 'cost of revenues': 'cogs',
          'cash and cash equivalents': 'cash_and_equivalents'}
LOOKUP = alias_lookup(ALIASES) | {header_key(k): v for k, v in LABELS.items()}
NUMBER = re.compile(r'(?<!\w)[−+-]?(?:\(\s*)?\d[\d,]*(?:\.\d+)?\s*\)?%?(?!\w)')


def propose(label):
    return LOOKUP.get(header_key(str(label)))


def tables(path, sheet=None):
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.reader(stream))
        yield None, rows, None
    elif path.suffix.lower() == '.json':
        def unique_object(pairs):
            result = {}
            for key, value in pairs:
                if key in result:
                    raise ModelError(f'Duplicate JSON key: {key}')
                result[key] = value
            return result
        value = json.loads(path.read_text(), parse_float=str, object_pairs_hook=unique_object)
        if not isinstance(value, list) or any(not isinstance(r, dict) for r in value):
            raise ModelError('JSON must be an array of objects')
        headers = list(dict.fromkeys(k for r in value for k in r))
        yield None, [headers] + [[r.get(k) for k in headers] for r in value], value
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            active = [ws for ws in wb if any(any(v is not None for v in r) for r in ws.iter_rows(values_only=True))]
            if sheet:
                if sheet not in wb.sheetnames:
                    raise ModelError('XLSX sheet not found')
                active = [wb[sheet]]
            elif len(active) > 1:
                raise ModelError('Multiple nonempty sheets: specify --sheet')
            for ws in active:
                yield ws.title, [list(r) for r in ws.iter_rows(values_only=True)], None
        finally:
            wb.close()


def extract_file(path, context=None, sheet=None):
    path = Path(path).resolve()
    if path.suffix.lower() not in {'.csv', '.json', '.xlsx', '.pdf'}:
        raise ModelError('Supported candidate inputs: CSV, JSON, XLSX, text PDF')
    context = dict(context or {})
    digest = hashlib.sha256(path.read_bytes()).hexdigest()

    def candidate(label, raw, locator, excerpt, headers, metadata=None, flags=None):
        data = context | (metadata or {})
        data = {k: v for k, v in data.items() if k in METADATA}
        year = str(data.get('fiscal_year', '')).upper().removeprefix('FY').strip()
        data['fiscal_year'] = int(year) if re.fullmatch(r'\d{4}', year) else None
        value = '' if raw is None else str(raw)
        return Candidate(str(path), digest, locator, excerpt, str(label), value,
                         metric=propose(label), headers=headers, extraction_flags=flags or [], **data)

    if path.suffix.lower() == '.pdf':
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                words = page.extract_words()
                text = page.extract_text() or ''
                if not words:
                    yield candidate('', '', {'page': page.page_number, 'bbox': [0, 0, page.width, page.height]}, '', [], flags=['OCR_REQUIRED'])
                    continue
                lines = []
                for word in words:
                    if not lines or abs(lines[-1][0]['top'] - word['top']) > 3:
                        lines.append([])
                    lines[-1].append(word)
                years = []
                for line in lines:
                    year_words = [(w, int(w['text'])) for w in line if re.fullmatch(r'20\d{2}|19\d{2}', w['text'])]
                    if len(year_words) >= 2:
                        years = year_words
                        continue
                    merged = []
                    i = 0
                    while i < len(line):
                        w = dict(line[i])
                        if w['text'] == '(' and i + 2 < len(line) and NUMBER.fullmatch(line[i+1]['text']) and line[i+2]['text'] == ')':
                            w['text'] = '(' + line[i+1]['text'] + ')'
                            w['x1'] = line[i+2]['x1']
                            i += 2
                        merged.append(w)
                        i += 1
                    numeric = [w for w in merged if NUMBER.fullmatch(w['text'])]
                    if not numeric:
                        continue
                    label_words = [w['text'] for w in line if w['x0'] < numeric[0]['x0']]
                    label = ' '.join(label_words)
                    if not label:
                        continue
                    for w in numeric:
                        meta = {}
                        if years:
                            meta['fiscal_year'] = min(years, key=lambda y: abs(y[0]['x1'] - w['x1']))[1]
                        # Entire page stays available to Jev; never quietly truncate context.
                        yield candidate(label, w['text'], {'page': page.page_number, 'bbox': [w['x0'], w['top'], w['x1'], w['bottom']]}, text, [str(y[1]) for y in years], meta)
        return
    from openpyxl.utils import get_column_letter
    for title, rows, objects in tables(path, sheet):
        if not rows:
            continue
        headers = [str(v or '') for v in rows[0]]
        physical_lines = {}
        if path.suffix.lower() == '.csv':
            with path.open(encoding='utf-8-sig', newline='') as stream:
                reader = csv.reader(stream, strict=True)
                start = 1
                for record, cells in enumerate(reader, 1):
                    physical_lines[record] = [start, reader.line_num]
                    start = reader.line_num + 1
                    if len(cells) != len(headers):
                        raise ModelError(f'Ragged CSV record {record}')

        canonical = [propose(h) for h in headers]
        year_columns = {i: int(re.sub(r'^FY\s*', '', h.upper())) for i, h in enumerate(headers) if re.fullmatch(r'(?:FY\s*)?\d{4}', h.upper())}
        for ri, row in enumerate(rows[1:], 2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            metadata = {k: row[i] for i, k in enumerate(canonical) if k in METADATA and i < len(row) and row[i] is not None}
            excerpt = json.dumps(dict(headers=headers, row=row), ensure_ascii=False, default=str)
            for ci, value in enumerate(row):
                if ci >= len(headers) or canonical[ci] in METADATA or (year_columns and ci not in year_columns):
                    continue
                label = str(row[0]) if year_columns else headers[ci]
                meta = metadata | ({'fiscal_year': year_columns[ci]} if year_columns else {})
                locator = {'row': ri, 'column': ci + 1, 'sheet': title, 'cell': f'{get_column_letter(ci+1)}{ri}'}
                if ri in physical_lines:
                    locator['physical_lines'] = physical_lines[ri]
                if objects is not None:
                    locator['json_path'] = f'$[{ri-2}][{json.dumps(headers[ci])}]'
                flags = []
                if canonical[ci] and canonical.count(canonical[ci]) > 1:
                    flags.append('AMBIGUOUS_HEADER')
                yield candidate(label, value, locator, excerpt, headers, meta, flags)


def extract_run(inputs, run, context=None, sheet=None):
    store = Store(run, create=True)
    count = 0
    for path in inputs:
        for candidate in extract_file(path, context, sheet):
            from .gate import inspect_candidate
            inspect_candidate(candidate)
            count += store.add(candidate)
    if not store.list():
        raise ModelError('No candidate values found')
    return {'run': str(store.root), 'new_candidates': count, 'total_candidates': len(store.list())}
