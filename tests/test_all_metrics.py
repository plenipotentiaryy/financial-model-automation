"""Full-catalog acceptance and input/Excel-guard parity, using a real engine."""

import json
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from finmodel.config import ALIASES, METRICS, METRIC_ORDER, TEMPLATE_MAP
from finmodel.provenance import source_records
from finmodel.recalculate import detect_engine
from finmodel.samples import FULL_DATA, FULL_EXPECTED, initialize
from finmodel.schemas import ModelError
from finmodel.service import build_company, configurations, load_dataset
from finmodel.validate import validate_company


def configuration():
    return deepcopy({"aliases": ALIASES, "metric_definitions": METRICS, "template_map": TEMPLATE_MAP})


def company_from_rows(tmp_path, rows):
    path = tmp_path / "financials.json"
    path.write_text(json.dumps(rows), encoding="utf-8")
    dataset = load_dataset([path], configuration())
    assert not dataset.issues
    return dataset.companies["complete"]


@pytest.fixture(scope="module")
def all_outputs(tmp_path_factory):
    if detect_engine() is None:
        pytest.skip("Real LibreOffice is required to verify all formula results")
    root = tmp_path_factory.mktemp("all_metrics")
    initialize(root)
    outputs = {}
    for extension in ("csv", "json", "xlsx"):
        config = configuration()
        company = load_dataset([root / f"samples/all_metrics.{extension}"], config).companies["complete"]
        target = root / f"{extension}.xlsx"
        manifest = build_company(company, root / "samples/template.xlsx", target, config, calculate=True)
        assert manifest["outcome"] == "SUCCESS", manifest["issues"]
        outputs[extension] = (load_workbook(target, data_only=True), load_workbook(target, data_only=False), manifest)
    return outputs


@pytest.mark.engine
@pytest.mark.parametrize("metric_id", METRIC_ORDER)
def test_every_metric_recalculates_for_all_formats(all_outputs, metric_id):
    metric = METRICS[metric_id]
    for values, formulas, manifest in all_outputs.values():
        sheet, address = metric["cell"].split("!")
        report_sheet, report_address = metric["report_cell"].split("!")
        assert values[sheet][address].value == pytest.approx(FULL_EXPECTED[metric_id], rel=1e-12, abs=1e-10)
        assert values[report_sheet][report_address].value == pytest.approx(FULL_EXPECTED[metric_id], rel=1e-12, abs=1e-10)
        assert formulas[sheet][address].value == metric["excel_formula"]
        assert manifest["metric_statuses"][metric_id]["status"] == "OK"
        assert manifest["calculation"]["status"] == "RECALCULATED"


@pytest.mark.parametrize("metric_id", METRIC_ORDER)
def test_each_metric_declares_missing_dependency(tmp_path, metric_id):
    rows = deepcopy(FULL_DATA)
    field, period = METRICS[metric_id]["requires"][0].split(":")
    year = {"previous": 2024, "latest": 2025, "start": 2023}[period]
    next(row for row in rows if row["fiscal_year"] == year)[field] = None
    company = company_from_rows(tmp_path, rows)
    status = company.metric_statuses[metric_id]
    assert status["status"] == "MISSING_INPUTS"
    assert field in status["message"] and str(year) in status["message"]


INVALID_SCENARIOS = [
    ([(2025, "revenue", 0)], ["ebitda_margin", "gross_margin", "ebit_margin", "net_profit_margin", "fcf_margin", "ev_to_revenue", "dso", "cash_conversion_cycle"]),
    ([(2025, "ebitda", 0)], ["net_debt_to_ebitda", "ev_to_ebitda", "cash_conversion", "net_leverage_trend"]),
    ([(2024, "ebitda", -2)], ["net_leverage_trend"]),
    ([(2025, "monthly_recurring_revenue", -1)], ["arr"]),
    ([(2025, "marginal_tax_rate", 1.2)], ["wacc"]),
    ([(2025, "market_value_debt", -1)], ["wacc"]),
    ([(2025, "market_cap", 0), (2025, "market_value_debt", 0)], ["wacc"]),
    ([(2025, "pre_tax_cost_of_debt", -0.01)], ["wacc"]),
    ([(2025, "capex", -1)], ["free_cash_flow", "fcf_margin"]),
    ([(2025, "current_liabilities", 0)], ["current_ratio", "quick_ratio"]),
    ([(2025, "inventory", 70)], ["quick_ratio"]),
    ([(2025, "interest_expense", 0)], ["interest_coverage"]),
    ([(2024, "total_assets", 0), (2025, "total_assets", 0)], ["roa"]),
    ([(2024, "total_equity", -40)], ["roe"]),
    ([(2025, "net_income", -1)], ["price_to_earnings"]),
    ([(2023, "revenue", 0)], ["revenue_cagr"]),
    ([(2023, "ebitda", -1)], ["ebitda_cagr"]),
    ([(2025, "cogs", 0)], ["dio", "dpo", "cash_conversion_cycle"]),
    ([(2024, "accounts_receivable", -1)], ["dso", "cash_conversion_cycle"]),
    ([(2025, "accounts_payable", -1)], ["dpo", "cash_conversion_cycle"]),
]


@pytest.mark.engine
@pytest.mark.skipif(detect_engine() is None, reason="LibreOffice unavailable")
@pytest.mark.parametrize("changes,affected", INVALID_SCENARIOS)
def test_invalid_values_have_matching_blank_excel_results(project, changes, affected):
    rows = deepcopy(FULL_DATA)
    for year, field, value in changes:
        next(row for row in rows if row["fiscal_year"] == year)[field] = value
    company = company_from_rows(project, rows)
    assert not company.blocked
    target = project / "edge.xlsx"
    manifest = build_company(company, project / "samples/template.xlsx", target, configuration(), calculate=True)
    assert manifest["outcome"] == "SUCCESS_WITH_WARNINGS", manifest["issues"]
    wb = load_workbook(target, data_only=True)
    for metric_id in affected:
        assert company.metric_statuses[metric_id]["status"] in {"INVALID_DENOMINATOR", "INVALID_INPUT", "INVALID_ASSUMPTION", "INCOMPATIBLE_PERIODS"}
        for key in ("cell", "report_cell"):
            sheet, address = METRICS[metric_id][key].split("!")
            assert wb[sheet][address].value in {None, ""}
    for ws in wb:
        assert not any(cell.data_type == "e" for row in ws for cell in row)


def test_cagr_uses_elapsed_years_and_never_substitutes_endpoints(tmp_path):
    rows = [deepcopy(FULL_DATA[0]), deepcopy(FULL_DATA[-1])]
    company = company_from_rows(tmp_path, rows)
    assert company.metric_statuses["revenue_cagr"]["status"] == "OK"
    assert company.metric_statuses["revenue_growth_yoy"]["status"] == "MISSING_INPUTS"
    rows[0]["revenue"] = None
    company = company_from_rows(tmp_path, rows)
    assert company.metric_statuses["revenue_cagr"]["status"] == "MISSING_INPUTS"
    company = company_from_rows(tmp_path, [deepcopy(FULL_DATA[-1])])
    assert company.metric_statuses["revenue_cagr"]["status"] == "INCOMPATIBLE_PERIODS"


def test_no_wacc_book_debt_fallback(tmp_path):
    rows = deepcopy(FULL_DATA)
    rows[-1].pop("market_value_debt")
    company = company_from_rows(tmp_path, rows)
    assert company.value("total_debt") is not None
    assert company.metric_statuses["wacc"]["status"] == "MISSING_INPUTS"


def test_all_dependency_lineage_and_cagr_multiple_destinations(tmp_path):
    company = company_from_rows(tmp_path, FULL_DATA)
    config = configuration()
    records = source_records(company, config["template_map"], config["metric_definitions"])
    for metric_id, metric in METRICS.items():
        for requirement in metric["requires"]:
            field, period = requirement.split(":")
            row = next(r for r in records if r["metric"] == field and r["fiscal_year"] == company.period_year(period))
            assert metric_id in row["dependent_metrics"].split(", ")
            assert row["destination_cell"] is not None
    two_years = company_from_rows(tmp_path, FULL_DATA[1:])
    records = source_records(two_years, config["template_map"], config["metric_definitions"])
    row = next(r for r in records if r["metric"] == "revenue" and r["fiscal_year"] == 2024)
    assert set(row["destination_cell"].split(", ")) == {"Inputs!B7", "Inputs!D7"}


def test_invalid_statuses_clear_on_revalidation(tmp_path):
    company = company_from_rows(tmp_path, FULL_DATA)
    observation = company.observations[("marginal_tax_rate", 2025)]
    from decimal import Decimal
    observation.canonical_value = Decimal("2")
    validate_company(company)
    assert company.metric_statuses["wacc"]["status"] == "INVALID_ASSUMPTION"
    observation.canonical_value = Decimal("0.2")
    validate_company(company)
    assert company.metric_statuses["wacc"]["status"] == "OK"
    assert not company.issues


def test_old_map_is_rejected_with_migration_instructions(tmp_path):
    import yaml
    mapping = deepcopy(TEMPLATE_MAP)
    mapping["version"] = 1
    path = tmp_path / "old.yaml"
    path.write_text(yaml.safe_dump(mapping))
    with pytest.raises(ModelError, match="v2 is required"):
        configurations(template_map=path)


@pytest.mark.parametrize("sheet,address,value", [("Data_Dictionary", "B1", 1), ("Calculations", "B6", None), ("Report", "B34", None)])
def test_incomplete_or_old_workbook_cannot_claim_full_coverage(project, sheet, address, value):
    template = project / "samples/template.xlsx"
    wb = load_workbook(template)
    wb[sheet][address] = value
    wb.save(template)
    company = company_from_rows(project, FULL_DATA)
    target = project / "bad_template.xlsx"
    manifest = build_company(company, template, target, configuration())
    assert manifest["outcome"] == "FAILED"
    assert not target.exists()


@pytest.mark.engine
@pytest.mark.skipif(detect_engine() is None, reason="LibreOffice unavailable")
def test_zero_inputs_and_negative_economically_valid_results(project):
    rows = deepcopy(FULL_DATA)
    rows[-1].update(monthly_recurring_revenue=0, capex=0, market_value_debt=0, cash_flow_from_operations=-2, accounts_payable=100)
    company = company_from_rows(project, rows)
    target = project / "valid_edges.xlsx"
    manifest = build_company(company, project / "samples/template.xlsx", target, configuration(), calculate=True)
    assert manifest["outcome"] == "SUCCESS", manifest["issues"]
    wb = load_workbook(target, data_only=True)
    def result(key):
        sheet, address = METRICS[key]["cell"].split("!")
        return wb[sheet][address].value
    assert result("arr") == 0
    assert result("free_cash_flow") == -2000000
    assert result("cash_conversion") == pytest.approx(-0.1)
    assert result("wacc") == pytest.approx(0.1)
    assert result("cash_conversion_cycle") < 0
