import json
from copy import deepcopy
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from finmodel.provenance import source_records
from finmodel.samples import ACME
from finmodel.schemas import ModelError
from finmodel.service import build_company, report_path
from finmodel.template import SHEETS, validate_map


def formulas(wb):
    return {(ws.title, c.coordinate): c.value for ws in wb for row in ws for c in row if c.data_type == "f"}


def build(project, company, config, **kwargs):
    path = project / "outputs/model.xlsx"
    manifest = build_company(company, project / "samples/template.xlsx", path, config, **kwargs)
    return path, manifest


def test_sample_model_inputs_and_formulas(project, make_company, config):
    path, manifest = build(project, make_company(), config)
    wb = load_workbook(path)
    template = load_workbook(project / "samples/template.xlsx")
    assert set(wb.sheetnames) == set(SHEETS)
    assert [wb["Inputs"][address].value for address in ["B7", "C7", "C8", "C9", "C10"]] == [80000000, 100000000, 20000000, 35000000, 5000000]
    assert formulas(wb) == formulas(template)
    assert "B4/Inputs!C8" in wb["Calculations"]["B5"].value
    assert wb.calculation.fullCalcOnLoad and wb.calculation.forceFullCalc
    assert manifest["calculation"]["status"] == "FORMULA_NOT_RECALCULATED"
    assert wb["Report"]["D11"].value == "MISSING_INPUTS"
    assert wb["Report"]["D12"].value == "MISSING_INPUTS"
    assert wb["Report"]["B11"].data_type == "f" and wb["Report"]["B12"].data_type == "f"
    assert load_workbook(path, data_only=True)["Calculations"]["B2"].value is None


def test_no_missing_to_zero_in_workbook(project, make_company, config):
    rows = deepcopy(ACME)
    rows[1][8] = ""
    path, _ = build(project, make_company(rows), config)
    assert load_workbook(path)["Inputs"]["C8"].value is None


def test_formula_style_and_unrelated_cell_preservation(project, make_company, config):
    template = project / "samples/template.xlsx"
    wb = load_workbook(template)
    wb["Inputs"]["D20"] = "retain me"
    wb["Inputs"]["D21"] = "=SUM(C9:C10)"
    wb["Inputs"]["C8"] = 999999
    wb.save(template)
    baseline = load_workbook(template)
    rows = deepcopy(ACME)
    rows[1][8] = ""
    path, _ = build(project, make_company(rows), config)
    actual = load_workbook(path)
    assert formulas(actual) == formulas(baseline)
    assert actual["Inputs"]["C8"].value is None
    assert actual["Inputs"]["D20"].value == "retain me"
    allowed = validate_map(baseline, config["template_map"])
    for ws in baseline:
        for row in ws:
            for cell in row:
                assert actual[ws.title][cell.coordinate]._style == cell._style
                writable = any(sheet == ws.title and bounds[0] <= cell.column <= bounds[2] and bounds[1] <= cell.row <= bounds[3] for sheet, bounds in allowed)
                if not writable:
                    assert actual[ws.title][cell.coordinate].value == cell.value
    assert actual["Inputs"].column_dimensions["B"].width == baseline["Inputs"].column_dimensions["B"].width


@pytest.mark.parametrize("change", ["formula", "overlap", "sheet", "merged", "capacity", "duplicate_field"])
def test_bad_template_mapping(project, make_company, config, change):
    mapping = config["template_map"]
    if change == "formula":
        mapping["financials"]["revenue"]["latest"] = "Calculations!B2"
    elif change == "overlap":
        mapping["financials"]["revenue"]["latest"] = mapping["financials"]["ebitda"]["latest"]
    elif change == "sheet":
        mapping["financials"]["revenue"]["latest"] = "NoSuchSheet!C7"
    elif change == "merged":
        template = project / "samples/template.xlsx"
        wb = load_workbook(template)
        wb["Inputs"].merge_cells("C7:D7")
        wb.save(template)
    elif change == "capacity":
        mapping["regions"]["source_map"] = "Source_Map!A5:Q5"
    elif change == "duplicate_field":
        mapping["assumptions"]["revenue"] = {"latest": "Assumptions!B12"}
    path, manifest = build(project, make_company(), config)
    assert manifest["outcome"] == "FAILED"
    assert not path.exists()
    assert report_path(path).is_dir()


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_formula_injection_is_literal_text_and_negative_numeric_is_numeric(project, make_company, config, prefix):
    rows = deepcopy(ACME)
    for row in rows:
        row[1] = prefix + 'HYPERLINK("https://example.invalid","test")'
    rows[1][8] = -20
    path, _ = build(project, make_company(rows), config)
    wb = load_workbook(path)
    assert wb["Inputs"]["B2"].data_type == "s"
    assert wb["Inputs"]["C8"].value == -20000000
    assert wb["Inputs"]["C8"].data_type == "n"
    rows = deepcopy(ACME)
    rows[1][8] = "=1+1"
    path2 = project / "bad.xlsx"
    result = build_company(make_company(rows), project / "samples/template.xlsx", path2, config)
    assert result["outcome"] == "FAILED" and not path2.exists()


def test_reports_have_real_lineage_and_hashes(project, make_company, config):
    path, manifest = build(project, make_company(), config)
    directory = report_path(path)
    assert {p.name for p in directory.iterdir()} == {"validation.json", "validation.txt", "source_map.csv", "run_manifest.json"}
    assert len(manifest["template"]["sha256"]) == 64
    assert len(manifest["input_files"][0]["sha256"]) == 64
    loaded = json.loads((directory / "run_manifest.json").read_text())
    revenue = next(o for o in loaded["observations"] if o["metric"] == "revenue" and o["fiscal_year"] == 2025)
    assert revenue["canonical_value"] == "100000000"
    assert revenue["sources"][0]["source_row"] == 3
    assert "FORMULA_NOT_RECALCULATED" in (directory / "validation.txt").read_text()


def test_no_overwrite_without_flag(project, make_company, config):
    path, _ = build(project, make_company(), config)
    baseline = path.read_bytes()
    with pytest.raises(ModelError, match="already exists"):
        build(project, make_company(), config)
    assert path.read_bytes() == baseline
    assert build(project, make_company(), config, overwrite=True)[1]["outcome"] != "FAILED"


def test_cannot_overwrite_input_or_template(project, make_company, config):
    template = project / "samples/template.xlsx"
    with pytest.raises(ModelError, match="overwrite its template"):
        build_company(make_company(), template, template, config, overwrite=True)


def test_preserve_unrelated_report_files(project, make_company, config):
    path, _ = build(project, make_company(), config)
    unrelated = report_path(path) / "my_notes.txt"
    unrelated.write_text("keep")
    baseline = path.read_bytes()
    with pytest.raises(ModelError, match="unrelated"):
        build(project, make_company(), config, overwrite=True)
    assert unrelated.read_text() == "keep" and path.read_bytes() == baseline


def test_excel_precision_loss_blocks_export(project, make_company, config):
    rows = deepcopy(ACME)
    rows[1][7] = "1234567890123456"
    path, manifest = build(project, make_company(rows), config)
    assert manifest["outcome"] == "FAILED" and not path.exists()
    assert "15-digit" in manifest["issues"][-2].message or any("15-digit" in i.message for i in manifest["issues"])


def test_reordered_config_does_not_shift_metric_statuses(project, make_company, config):
    config["metric_definitions"] = dict(reversed(list(config["metric_definitions"].items())))
    path, _ = build(project, make_company(), config)
    wb = load_workbook(path)
    assert wb["Report"]["D7"].value == "OK"
    assert wb["Report"]["D11"].value == "MISSING_INPUTS"


def test_xlsx_date_in_financial_field_reports_invalid_number(project, config):
    from datetime import datetime
    from finmodel.service import load_dataset

    source = project / "samples/acme.xlsx"
    wb = load_workbook(source)
    wb["Financials"]["H3"] = datetime(2025, 1, 1)
    wb.save(source)
    company = load_dataset([source], config).companies["acme"]
    path, manifest = build(project, company, config)
    assert manifest["outcome"] == "FAILED" and not path.exists()
    assert "2025-01-01" in (report_path(path) / "run_manifest.json").read_text()


def test_custom_template_default_styles_are_not_changed(project, make_company, config):
    from openpyxl.styles.cell_style import StyleArray

    template = project / "samples/template.xlsx"
    wb = load_workbook(template)
    wb["Report"]["D7"]._style = StyleArray()
    wb["Report"].row_dimensions[7].height = 25
    wb.save(template)
    path, _ = build(project, make_company(), config)
    saved = load_workbook(path)
    assert not saved["Report"]["D7"].has_style
    assert saved["Report"].row_dimensions[7].height == 25


def test_output_symlink_is_not_followed(project, make_company, config):
    target = project / "unrelated.xlsx"
    target.write_bytes(b"keep me")
    output = project / "link.xlsx"
    output.symlink_to(target)
    with pytest.raises(ModelError, match="symbolic link"):
        build_company(make_company(), project / "samples/template.xlsx", output, config, overwrite=True)
    assert target.read_bytes() == b"keep me"
