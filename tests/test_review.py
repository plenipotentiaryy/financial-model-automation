"""Offline fixtures only. MockTransport responses are NOT live Jev evidence."""
import csv
import hashlib
import json
from decimal import Decimal

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from finmodel.review.contracts import CHECKS, MODEL
from finmodel.review.extract import extract_run, extract_file
from finmodel.review.store import Store
from finmodel.review.gate import decide, export_companies, inspect_candidate, snapshot
from finmodel.review.jev import JevProvider, verify_run
from finmodel.review.ui import create_app
from finmodel.schemas import ModelError

CTX = dict(company_id='acme', company_name='Acme', fiscal_year=2025, period_type='FY', currency='USD', monetary_unit='millions', scope='consolidated')


def fixture_response(score=.999):
    return {'model': MODEL, 'usage': {'input_tokens': 1000, 'output_tokens': 80}, 'answers': {
        k: {'type': 'choice', 'choice': 'match', 'probabilities': {'match': score, 'other_metric': (1-score)/2, 'insufficient_context': (1-score)/2}} for k in CHECKS}}


def fixture_provider(store, handler=None, budget='0.1'):
    client = httpx.Client(transport=httpx.MockTransport(handler or (lambda req: httpx.Response(200, json=fixture_response()))))
    return JevProvider(store, budget, client=client, key='offline-test-fixture', sleeper=lambda _: None)


def run(tmp_path, rows=None):
    path = tmp_path / 'input.json'
    path.write_text(json.dumps(rows or [CTX | {'revenue': 100, 'ebitda': 20}]))
    extract_run([path], tmp_path / 'run')
    return Store(tmp_path / 'run')


def test_four_formats_and_real_coordinates(tmp_path):
    row = CTX | {'revenue': '-1,000', 'ebitda': '0', 'ORRR': 55, 'Gross margin': 60}
    jsonfile = tmp_path/'a.json'; jsonfile.write_text(json.dumps([row]))
    csvfile = tmp_path/'a.csv'
    with csvfile.open('w') as f:
        writer=csv.DictWriter(f,fieldnames=list(row));writer.writeheader();writer.writerow(row)
    wb=Workbook();wb.active.title='Financials';wb.active.append(list(row));wb.active.append(list(row.values()));wb.save(tmp_path/'a.xlsx')
    all_candidates = [list(extract_file(p)) for p in (jsonfile,csvfile,tmp_path/'a.xlsx')]
    for candidates in all_candidates:
        revenue=next(c for c in candidates if c.metric=='revenue')
        assert not inspect_candidate(revenue)
        assert revenue.normalized_value=='-1000000000'
        assert revenue.locator['row']==2 and revenue.locator['column']==8
        assert {c.raw_label for c in candidates if c.metric is None}=={'ORRR','Gross margin'}
    assert all_candidates[0][0].locator['json_path']=='$[0]["revenue"]'
    assert all_candidates[2][0].locator['sheet']=='Financials'
    assert all_candidates[2][0].locator['cell']=='H2'


def write_pdf(path, lines):
    """Tiny real PDF test document; no PDF-library fixture responses."""
    content='BT /F1 12 Tf ' + ' '.join(f'1 0 0 1 {x} {y} Tm ({text.replace(chr(40),chr(92)+chr(40)).replace(chr(41),chr(92)+chr(41))}) Tj' for x,y,text in lines) + ' ET'
    objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',f'<< /Length {len(content.encode())} >>\nstream\n{content}\nendstream'.encode()]
    pdf=b'%PDF-1.4\n'; offsets=[0]
    for i,obj in enumerate(objects,1):
        offsets.append(len(pdf));pdf+=f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n'
    xref=len(pdf);pdf+=f'xref\n0 6\n0000000000 65535 f \n'.encode()+b''.join(f'{o:010d} 00000 n \n'.encode() for o in offsets[1:])+f'trailer << /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode();path.write_bytes(pdf)


def test_pdf_year_columns_negatives_and_scan(tmp_path):
    path=tmp_path/'table.pdf'
    write_pdf(path,[(30,760,'Acme USD millions consolidated FY'),(280,720,'2024'),(400,720,'2025'),(30,690,'Revenue'),(280,690,'80'),(400,690,'100'),(30,660,'EBITDA'),(280,660,'(5)'),(400,660,'20')])
    candidates=list(extract_file(path,CTX))
    assert len(candidates)==4
    for c in candidates: inspect_candidate(c)
    assert {(c.metric,c.fiscal_year,c.normalized_value) for c in candidates}=={('revenue',2024,'80000000'),('revenue',2025,'100000000'),('ebitda',2024,'-5000000'),('ebitda',2025,'20000000')}
    assert all(c.locator['page']==1 and len(c.locator['bbox'])==4 for c in candidates)
    scan=tmp_path/'scan.pdf';write_pdf(scan,[])
    assert 'OCR_REQUIRED' in inspect_candidate(next(extract_file(scan,CTX)))


def test_year_columns_unknowns_and_lost_context(tmp_path):
    path=tmp_path/'table.csv';path.write_text('metric,2024,2025\nrevenue,80,100\nORRR,5,6\nDepreciation amortization and other,1,2\n')
    cs=list(extract_file(path,CTX));assert len(cs)==6
    assert cs[0].fiscal_year==2024 and cs[1].fiscal_year==2025
    assert 'MIXED_DA_NOT_SUPPORTED' in inspect_candidate(cs[-1])
    assert 'MISSING_COMPANY' in inspect_candidate(next(extract_file(path)))


@pytest.mark.parametrize('raw,unit,expected', [('1.2','millions','1200000.0'),('0','thousands','0'),('(1,200)','units','-1200')])
def test_normalization(tmp_path,raw,unit,expected):
    store=run(tmp_path,[CTX|{'monetary_unit':unit,'revenue':raw}]);c=store.list()[0]['candidate']
    assert not inspect_candidate(c) and c.normalized_value==expected


@pytest.mark.parametrize('raw', ['1,20','=1+2','12%','NaN','unavailable',''])
def test_invalid_or_unavailable_never_approved(tmp_path,raw):
    store=run(tmp_path,[CTX|{'revenue':raw}]);c=store.list()[0]['candidate']
    assert inspect_candidate(c)
    with pytest.raises(ModelError):decide(store,c.candidate_id,1,'accept','reviewed')


def test_high_scores_cache_and_export(tmp_path):
    store=run(tmp_path);calls=[]
    def handler(req):
        calls.append(req);assert req.url==httpx.URL('https://api.typesafe.ai/v1/systemone');return httpx.Response(200,json=fixture_response())
    result=verify_run(store,'1',fixture_provider(store,handler))
    assert result['accepted_auto']==2 and len(calls)==2
    company=export_companies(store)['acme'];assert company.value('revenue')==Decimal('100000000')
    assert company.review_audit['history']
    result=verify_run(store,'1',fixture_provider(store,handler));assert len(calls)==2
    assert all(r['cached'] for r in result['results'])
    assert company.observations[('revenue',2025)].sources[0].source_sha256


@pytest.mark.parametrize('response', [fixture_response(.8), {'model':MODEL}, fixture_response()|{'model':'jev-latest'}, fixture_response()|{'answers':{}}, [], fixture_response()|{'usage':[]}])
def test_abstention_incomplete_and_version(tmp_path,response):
    store=run(tmp_path)
    result=verify_run(store,'1',fixture_provider(store,lambda req:httpx.Response(200,json=response)))
    assert result['accepted_auto']==0
    with pytest.raises(ModelError):export_companies(store)


def test_error_retries_budget_and_missing_key(tmp_path,monkeypatch):
    store=run(tmp_path);calls=[]
    def unavailable(req):calls.append(1);return httpx.Response(529)
    result=verify_run(store,'1',fixture_provider(store,unavailable))
    assert len(calls)==6 and result['accepted_auto']==0
    calls.clear();result=verify_run(store,'0.00001',fixture_provider(store,unavailable,'0.00001'))
    assert not calls and all('COST_LIMIT' in r['reason'] for r in result['results'])
    monkeypatch.delenv('TYPESAFE_API_KEY',raising=False)
    result=verify_run(store,'1');assert result['status']=='NOT_RUN'


def test_retry_success_and_context_limit(tmp_path):
    store=run(tmp_path,[CTX|{'revenue':100}]);calls=[]
    def handler(req):calls.append(1);return httpx.Response(429) if len(calls)==1 else httpx.Response(200,json=fixture_response())
    assert verify_run(store,'1',fixture_provider(store,handler))['accepted_auto']==1
    assert len(calls)==2
    c=store.list()[0]['candidate'];c.excerpt='x'*25000
    assert 'CONTEXT_LIMIT' in fixture_provider(store).evaluate(c).reason


def test_conflicts_selection_revision_and_source_tamper(tmp_path):
    store=run(tmp_path,[CTX|{'revenue':100},CTX|{'revenue':200}]);cs=[r['candidate'] for r in store.list()]
    assert verify_run(store,'1',fixture_provider(store))['accepted_auto']==0
    with pytest.raises(ModelError):decide(store,cs[0].candidate_id,1,'accept','pick')
    decide(store,cs[0].candidate_id,1,'select','correct row after reviewing evidence')
    assert len(export_companies(store)['acme'].observations)==1
    decide(store,cs[0].candidate_id,1,'correct','fix field',{'metric':'gross_profit'})
    with pytest.raises(ModelError):export_companies(store)
    with pytest.raises(ModelError):decide(store,cs[0].candidate_id,1,'accept','stale')
    decide(store,cs[0].candidate_id,2,'accept','reviewed revised mapping')
    assert Store(store.root).get(cs[0].candidate_id)['state']=='accepted_manual'
    (tmp_path/'input.json').write_text('[]')
    with pytest.raises(ModelError,match='SOURCE_CHANGED'):export_companies(store)
    with pytest.raises(ModelError,match='SOURCE_CHANGED'):decide(store,cs[0].candidate_id,2,'accept','cannot bypass')


def test_ambiguous_headers_and_idempotence(tmp_path):
    p=tmp_path/'a.csv';p.write_text('company_id,company_name,fiscal_year,period_type,scope,currency,monetary_unit,revenue,sales\na,A,2025,FY,consolidated,USD,units,100,100\n')
    first=extract_run([p],tmp_path/'run');second=extract_run([p],tmp_path/'run')
    assert first['new_candidates']==2 and second['new_candidates']==0
    store=Store(tmp_path/'run');assert verify_run(store,'1',fixture_provider(store))['accepted_auto']==0


def test_ui_origin_token_persistence_evidence_and_coverage(tmp_path):
    store=run(tmp_path);c=store.list()[0]['candidate']
    client=TestClient(create_app(store,token='secret'),base_url='http://127.0.0.1:8765')
    assert client.get('/').status_code==200
    snapshot=client.get('/api/run').json();assert snapshot['candidates'][0]['excerpt']
    body=dict(candidate_id=c.candidate_id,revision=1,action='accept',reason='Reviewed source')
    assert client.post('/api/decision',json=body).status_code==403
    assert client.post('/api/decision',json=body,headers={'Origin':'https://evil.example','X-Review-Token':'secret'}).status_code==403
    headers={'Origin':'http://127.0.0.1:8765','X-Review-Token':'secret'}
    assert client.post('/api/decision',json=body,headers=headers).status_code==200
    assert 'acme' in client.get('/api/run').json()['coverage']
    assert client.get('/api/source/'+c.candidate_id).status_code==200
    assert client.get('/api/run',headers={'Host':'evil.example'}).status_code==403
    restarted=TestClient(create_app(Store(store.root),token='new'),base_url='http://127.0.0.1:8765')
    assert restarted.get('/api/run').json()['candidates'][0]['state']=='accepted_manual'
    body['action']='reject'
    assert restarted.post('/api/decision',json=body,headers=headers|{'X-Review-Token':'new'}).status_code==200
    assert 'blocked' in restarted.get('/api/run').json()['coverage']


def test_unapproved_cannot_leak_to_excel(tmp_path):
    from finmodel.cli import main
    from finmodel.samples import initialize
    initialize(tmp_path/'demo')
    store=run(tmp_path)
    out=tmp_path/'model.xlsx';template=tmp_path/'demo/samples/template.xlsx'
    assert main(['build','--run',str(store.root),'--template',str(template),'--output',str(out)])==1
    assert not out.exists()
    c=next(r['candidate'] for r in store.list() if r['candidate'].metric=='revenue')
    decide(store,c.candidate_id,1,'accept','Reviewed source')
    assert main(['build','--run',str(store.root),'--template',str(template),'--output',str(out)])==0
    wb=load_workbook(out);assert wb['Inputs']['C7'].value==100000000 and wb['Inputs']['C8'].value is None
    manifest=json.loads((tmp_path/'model.reports/run_manifest.json').read_text())
    assert manifest['input_mode']=='approved_candidates' and manifest['review_audit_sha256']
    assert main(['build','--run',str(store.root),'--template',str(template),'--output',str(out),'--overwrite'])==0


def test_exact_json_lexemes_and_duplicate_keys(tmp_path):
    p=tmp_path/'exact.json'
    p.write_text('['+json.dumps(CTX)[:-1]+',"revenue":0.123456789012345678901}]')
    c=next(extract_file(p));assert c.raw_value=='0.123456789012345678901'
    inspect_candidate(c);assert Decimal(c.normalized_value)==Decimal('123456.789012345678901000000')
    p.write_text('[{"revenue":1,"revenue":2}]')
    with pytest.raises(ModelError,match='Duplicate JSON key'):list(extract_file(p))


def test_identical_files_keep_both_sources_and_do_not_conflict(tmp_path):
    a=tmp_path/'a.json';b=tmp_path/'b.json';a.write_text(json.dumps([CTX|{'revenue':100}]));b.write_bytes(a.read_bytes())
    extract_run([a,b],tmp_path/'run');store=Store(tmp_path/'run')
    assert len(store.list())==2
    verify_run(store,'1',fixture_provider(store))
    assert len(export_companies(store)['acme'].observations[('revenue',2025)].sources)==2


def test_human_decision_wins_during_provider_call(tmp_path):
    store=run(tmp_path,[CTX|{'revenue':100}]);c=store.list()[0]['candidate']
    def response(req):
        decide(store,c.candidate_id,1,'reject','human rejected while network call in flight')
        return httpx.Response(200,json=fixture_response())
    verify_run(store,'1',fixture_provider(store,response))
    assert store.get(c.candidate_id)['state']=='rejected'


def test_currency_and_scope_conflicts_never_auto_approve(tmp_path):
    store=run(tmp_path,[CTX|{'revenue':100},CTX|{'ebitda':20,'currency':'PLN','scope':'standalone'}])
    assert verify_run(store,'1',fixture_provider(store))['accepted_auto']==0


def test_ragged_csv_and_physical_coordinates(tmp_path):
    p=tmp_path/'rows.csv';p.write_text('label,2025\n"two\nlines",42\nrevenue,100\n')
    cs=list(extract_file(p,CTX));assert cs[0].locator['physical_lines']==[2,3] and cs[1].locator['physical_lines']==[4,4]
    p.write_text('label,2025\nrevenue,1,2\n')
    with pytest.raises(ModelError,match='Ragged'):list(extract_file(p,CTX))


def test_review_hash_rechecked_after_slow_engine(tmp_path,monkeypatch):
    from finmodel import service
    from finmodel.samples import initialize
    initialize(tmp_path/'demo');store=run(tmp_path,[CTX|{'revenue':100}]);c=store.list()[0]['candidate']
    decide(store,c.candidate_id,1,'accept','checked')
    company=export_companies(store)['acme']
    def changed_source(path):
        (tmp_path/'input.json').write_text('[]')
        return {'status':'RECALCULATED','engine':'test fixture'}
    monkeypatch.setattr(service,'recalculate',changed_source)
    out=tmp_path/'model.xlsx'
    manifest=service.build_company(company,tmp_path/'demo/samples/template.xlsx',out,service.configurations(),calculate=True)
    assert manifest['outcome']=='FAILED' and not out.exists()
    assert any('SOURCE_CHANGED' in issue.message for issue in manifest['issues'])


def test_selected_company_audit_isolation(tmp_path):
    store=run(tmp_path,[CTX|{'revenue':100},CTX|{'company_id':'other','company_name':'Other','revenue':300}])
    verify_run(store,'1',fixture_provider(store))
    company=export_companies(store,company_id='acme')['acme']
    assert all(c['company_id']=='acme' for c in company.review_audit['candidates'])
    assert all(e['payload'].get('company_id','acme')=='acme' for e in company.review_audit['history'])


@pytest.mark.parametrize('cost', ['zero', 'NaN', '-1', '0'])
def test_invalid_budget_is_argument_error(cost):
    from finmodel.cli import main
    with pytest.raises(SystemExit) as exc:
        main(['verify', '--run', 'irrelevant', '--provider', 'jev', '--max-cost-usd', cost])
    assert exc.value.code == 2
