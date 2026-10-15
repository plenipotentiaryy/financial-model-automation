import json
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from finmodel.cli import main
from finmodel.samples import ACME, initialize, write_csv
from finmodel.service import safe_name


def test_cli_workflow(project):
    output = project / "output.xlsx"
    assert main(["validate", "--input", str(project / "samples/acme.csv"), "--output-dir", str(project / "validation")]) == 0
    assert main(["build", "--input", str(project / "samples/acme.json"), "--template", str(project / "samples/template.xlsx"), "--output", str(output)]) == 0
    assert output.exists()
    assert main(["explain", "revenue_growth_yoy"]) == 0
    assert main(["explain", "unknown"]) == 1


def test_init_does_not_overwrite(project):
    template = project / "samples/template.xlsx"
    baseline = template.read_bytes()
    assert main(["init", "--directory", str(project)]) == 1
    assert template.read_bytes() == baseline
    assert main(["init", "--directory", str(project), "--overwrite"]) == 0


def test_batch_isolation_and_partial_failure(project):
    inputs = project / "samples/companies"
    invalid = deepcopy(ACME)
    for row in invalid:
        row[0], row[1] = "bad", "Bad Example"
    invalid[1][8] = "invalid number"
    write_csv(inputs / "bad.csv", invalid)
    (inputs / "malformed.json").write_text("{bad}")
    output = project / "batch"
    assert main(["batch", "--input-dir", str(inputs), "--template", str(project / "samples/template.xlsx"), "--output-dir", str(output)]) == 1
    books = list(output.glob("*.xlsx"))
    assert len(books) == 2
    identities = {}
    for path in books:
        wb = load_workbook(path)
        identities[wb["Inputs"]["B5"].value] = (wb["Inputs"]["B2"].value, wb["Inputs"]["C7"].value, wb["Inputs"]["C4"].value)
    assert identities == {"acme": ("Acme Example", 100000000, "PLN"), "beta": ("Beta Example", 300000, "EUR")}
    summary = json.loads((output / "batch_summary.json").read_text())
    assert summary["outcome"] == "FAILED"
    assert {entry["company_id"] for entry in summary["companies"] if entry["outcome"] == "FAILED"} == {"bad"}
    assert summary["input_issues"][0]["severity"] == "error"


def test_batch_conflict_across_files(project):
    inputs = project / "samples/companies"
    duplicate = deepcopy(ACME)
    duplicate[1][7] = 101
    write_csv(inputs / "conflicting.csv", duplicate)
    output = project / "batch"
    assert main(["batch", "--input-dir", str(inputs), "--template", str(project / "samples/template.xlsx"), "--output-dir", str(output)]) == 1
    books = list(output.glob("*.xlsx"))
    assert len(books) == 1
    assert load_workbook(books[0])["Inputs"]["B5"].value == "beta"


def test_multicompany_build_requires_selection(project):
    path = project / "multi.csv"
    other = deepcopy(ACME)
    for row in other:
        row[0], row[1] = "other", "Other Example"
    write_csv(path, ACME + other)
    args = ["build", "--input", str(path), "--template", str(project / "samples/template.xlsx"), "--output", str(project / "output.xlsx")]
    assert main(args) == 1
    assert main([*args, "--company", "other"]) == 0
    assert load_workbook(project / "output.xlsx")["Inputs"]["B5"].value == "other"


def test_malformed_input_validation_report(project):
    path = project / "bad.json"
    path.write_text("{broken")
    output = project / "validation"
    assert main(["validate", "--input", str(path), "--output-dir", str(output)]) == 1
    assert json.loads((output / "validation_summary.json").read_text())["input_issues"][0]["code"] == "INPUT_ERROR"


def test_invalid_arguments_exit_two():
    with pytest.raises(SystemExit) as exc:
        main(["build"])
    assert exc.value.code == 2


def test_malformed_build_preserves_error_without_inventing_identity(project):
    path = project / "bad.json"
    path.write_text("{broken")
    output = project / "unresolved.xlsx"
    assert main(["build", "--input", str(path), "--output", str(output)]) == 1
    manifest = json.loads((project / "unresolved.reports/run_manifest.json").read_text())
    assert manifest["company_id"] is None and manifest["output_file"] is None
    assert manifest["issues"][0]["code"] == "INPUT_ERROR"
    assert not output.exists()
