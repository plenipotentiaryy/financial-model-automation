import json
from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from finmodel.ingest import read_file
from finmodel.mapping import alias_lookup
from finmodel.normalize import parse_number
from finmodel.provenance import source_records
from finmodel.samples import ACME, HEADERS
from finmodel.schemas import ModelError
from finmodel.service import load_dataset


def test_normalize_millions(make_company):
    company = make_company()
    assert company.value("revenue") == Decimal("100000000")
    assert company.value("revenue", 2024) == Decimal("80000000")
    assert company.observations[("total_debt", 2025)].measurement_basis == "fiscal_year_end"
    assert company.observations[("ebitda", 2025)].measurement_basis == "annual_flow"


@pytest.mark.parametrize("raw,state", [(None, "missing"), ("", "missing"), ("missing", "missing"), ("not applicable", "not_applicable"), ("unavailable", "unavailable"), ("invalid", "invalid")])
def test_missing_states(raw, state):
    assert parse_number(raw) == (None, state)


@pytest.mark.parametrize("raw", ["nan", "Infinity", "1,000", "=1+1", True, {}, "100 PLN", "1e999"])
def test_invalid_number(raw):
    with pytest.raises(ModelError):
        parse_number(raw)


def test_no_missing_to_zero(make_company):
    rows = deepcopy(ACME)
    rows[1][8] = ""
    company = make_company(rows)
    assert company.value("ebitda") is None
    assert company.observations[("ebitda", 2025)].value_status == "missing"
    assert company.metric_statuses["ebitda_margin"]["status"] == "MISSING_INPUTS"
    rows[1][8] = 0
    assert make_company(rows).value("ebitda") == 0


@pytest.mark.parametrize("index,value,code", [(5, "EUR", "CURRENCY_MISMATCH"), (4, "standalone", "SCOPE_MISMATCH"), (3, "TTM", "INCOMPATIBLE_PERIODS"), (3, "Q1", "INCOMPATIBLE_PERIODS"), (6, "billions", "INVALID_UNIT"), (2, "2025.5", "INVALID_PERIOD"), (1, "Other company", "INVALID_IDENTITY"), (8, "broken", "INVALID_NUMBER")])
def test_conflicting_or_invalid_data(make_company, index, value, code):
    rows = deepcopy(ACME)
    rows[1][index] = value
    company = make_company(rows)
    assert company.blocked
    assert code in [i.code for i in company.issues]


def test_duplicate_conflict(make_company):
    rows = deepcopy(ACME)
    rows.append(list(rows[1]))
    rows[-1][7] = 110
    company = make_company(rows)
    assert company.blocked
    assert any(i.code == "AMBIGUOUS_SOURCE" for i in company.issues)


def test_identical_duplicates_keep_all_sources(make_company, config):
    rows = deepcopy(ACME)
    rows.append(list(rows[1]))
    company = make_company(rows)
    assert not company.blocked
    assert len(company.observations[("revenue", 2025)].sources) == 2
    sources = source_records(company, config["template_map"], config["metric_definitions"])
    revenue = [r for r in sources if r["metric"] == "revenue" and r["fiscal_year"] == 2025]
    assert [r["source_row"] for r in revenue] == [3, 4]


def test_equal_values_different_scales_keep_original_units(make_company, config):
    rows = deepcopy(ACME)
    duplicate = list(rows[1])
    duplicate[6] = "units"
    duplicate[7:] = [100000000, 20000000, 35000000, 5000000]
    rows.append(duplicate)
    company = make_company(rows)
    assert not company.blocked
    sources = source_records(company, config["template_map"], config["metric_definitions"])
    assert {r["original_unit"] for r in sources if r["metric"] == "revenue"} == {"millions", "units"}


def test_previous_year_is_strictly_yoy(make_company):
    rows = deepcopy(ACME)
    rows[0][2] = 2023
    company = make_company(rows)
    assert company.metric_statuses["revenue_growth_yoy"]["status"] == "MISSING_INPUTS"
    assert "FY2024" in company.metric_statuses["revenue_growth_yoy"]["message"]


@pytest.mark.parametrize("index,value,metric", [(7, 0, "ebitda_margin"), (8, 0, "net_debt_to_ebitda"), (8, -2, "net_debt_to_ebitda")])
def test_denominator_validation(make_company, index, value, metric):
    rows = deepcopy(ACME)
    rows[1][index] = value
    company = make_company(rows)
    assert not company.blocked
    assert company.metric_statuses[metric]["status"] == "INVALID_DENOMINATOR"


def test_previous_zero_denominator(make_company):
    rows = deepcopy(ACME)
    rows[0][7] = 0
    assert make_company(rows).metric_statuses["revenue_growth_yoy"]["status"] == "INVALID_DENOMINATOR"


def test_no_arr_or_wacc_guess(make_company):
    company = make_company()
    assert company.metric_statuses["arr"]["status"] == "MISSING_INPUTS"
    assert company.metric_statuses["wacc"]["status"] == "MISSING_INPUTS"
    assert "market_value_debt" in company.metric_statuses["wacc"]["message"]


def test_csv_json_xlsx_parity_and_provenance(project, config):
    companies = [load_dataset([project / f"samples/acme.{extension}"], config).companies["acme"] for extension in ["csv", "json", "xlsx"]]
    for company in companies:
        assert company.value("revenue") == Decimal("100000000")
        assert company.value("ebitda", 2024) is None
    assert len({tuple((k, o.canonical_value, o.value_status) for k, o in c.observations.items()) for c in companies}) == 1
    csv_source, json_source, xlsx_source = [c.observations[("revenue", 2025)].sources[0] for c in companies]
    assert csv_source.source_row == 3 and csv_source.source_column == "revenue"
    assert json_source.source_locator == '$[1]["revenue"]'
    assert xlsx_source.source_sheet == "Financials" and xlsx_source.source_locator == "Financials!H3"
    records = source_records(companies[0], config["template_map"], config["metric_definitions"])
    assert len(records) == 8
    revenue = next(r for r in records if r["metric"] == "revenue" and r["fiscal_year"] == 2025)
    assert revenue["destination_cell"] == "Inputs!C7"
    assert "ebitda_margin" in revenue["dependent_metrics"]


def test_xlsx_formula_input_is_rejected(project, config):
    path = project / "samples/acme.xlsx"
    wb = load_workbook(path)
    wb["Financials"]["H3"] = "=99+1"
    wb.save(path)
    assert load_dataset([path], config).companies["acme"].blocked


def test_multiple_xlsx_sheets_require_choice(project, config):
    path = project / "samples/acme.xlsx"
    wb = load_workbook(path)
    wb.create_sheet("Other")["A1"] = "data"
    wb.save(path)
    with pytest.raises(ModelError, match="--sheet"):
        read_file(path, config["aliases"])
    assert len(read_file(path, config["aliases"], "Financials")[0]) == 2


def test_alias_mapping_and_ambiguity(tmp_path, config):
    path = tmp_path / "aliases.csv"
    text = ",".join(HEADERS).replace("revenue", "Sales") + "\n" + ",".join(str(v) if v is not None else "" for v in ACME[1]) + "\n"
    path.write_text(text)
    assert load_dataset([path], config).companies["acme"].value("revenue") == Decimal("100000000")
    config["aliases"]["ebitda"].append("sales")
    with pytest.raises(ModelError, match="Ambiguous alias"):
        alias_lookup(config["aliases"])


def test_ambiguous_columns_fail(tmp_path, config):
    path = tmp_path / "bad.csv"
    path.write_text(",".join(HEADERS + ["sales"]) + "\n" + ",".join(str(v) if v is not None else "" for v in [*ACME[1], 10]))
    dataset = load_dataset([path], config)
    assert not dataset.companies
    assert dataset.issues[0].code == "INPUT_ERROR"


def test_json_exact_decimal_and_rate_scale(tmp_path, config):
    path = tmp_path / "precise.json"
    obj = dict(zip(HEADERS, ACME[1]))
    obj["revenue"] = "1.234567890123456789"
    obj["risk_free_rate"] = "0.04"
    obj["share_price"] = "12.50"
    path.write_text(json.dumps([obj]))
    company = load_dataset([path], config).companies["acme"]
    assert company.value("revenue") == Decimal("1234567.890123456789")
    assert company.value("risk_free_rate") == Decimal("0.04")
    assert company.value("share_price") == Decimal("12.50")


def test_duplicate_json_keys_fail(tmp_path, config):
    path = tmp_path / "bad.json"
    path.write_text('[{"company_id":"a","company_id":"b"}]')
    with pytest.raises(ModelError, match="Duplicate JSON"):
        read_file(path, config["aliases"])


def test_empty_input_has_issue(tmp_path, config):
    path = tmp_path / "empty.csv"
    path.write_text("")
    dataset = load_dataset([path], config)
    assert not dataset.companies
    assert any(i.code == "INPUT_ERROR" for i in dataset.issues)


def test_nonstandard_json_constants_fail(tmp_path, config):
    path = tmp_path / "bad.json"
    path.write_text('[{"revenue":NaN}]')
    with pytest.raises(ModelError, match="Nonstandard JSON"):
        read_file(path, config["aliases"])
